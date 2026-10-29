"""Planilha de entrada do Planejador (.xlsx) na identidade Bioz Green.

`gerar_modelo_xlsx()`            -> modelo em branco (1 linha de EXEMPLO por aba; exemplos são ignorados pelo sistema)
`gerar_modelo_xlsx(dados=Dados)` -> mesmo layout já preenchido com os dados carregados (para começar de onde você está)

Só a aba BOM (estruturas) tem origem no Tiny; todo o resto é digitado/colado por você.
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VERDE_ESCURO = "0B592F"
VERDE_CLARO = "8CA70A"
AMARELO = "EEB627"
LARANJA = "DF6F3B"
CREME = "F0E7D7"
CINZA = "595959"

_borda = Side(style="thin", color="D9D9D9")
BORDA = Border(top=_borda, bottom=_borda, left=_borda, right=_borda)
LOCAIS = "Plenno,LiveMax,Bioz"
LINHAS_ENTRADA = 300          # linhas pré-formatadas em creme para digitar

# tipo de coluna -> (formato numérico, largura)
FMT = {"txt": (None, 18), "desc": (None, 46), "int": ("#,##0", 14), "num": ("#,##0.##", 14),
       "data": ("DD/MM/YYYY", 15), "brl": ('"R$" #,##0.00', 14)}

# Cada aba: (nome, cor da aba, quem preenche, frequência, o que colocar, colunas[(nome, tipo, ajuda)], exemplo)
ABAS = [
    {
        "nome": "SKUs", "cor": VERDE_ESCURO, "quem": "Você", "freq": "Todo mês",
        "texto": "Um produto acabado por linha: fábrica, caixas, pallet, estoque, em produção e a meta do mês. "
                 "Estoque e Em produção em UNIDADES; reservado em CAIXAS.",
        "colunas": [
            ("sku", "txt", "Código do produto acabado (igual ao do Tiny)."),
            ("descricao", "desc", "Descrição do produto."),
            ("fabrica", "txt", "Onde é fabricado: Plenno, LiveMax ou Bioz (igual à aba Capacidade)."),
            ("un_por_cx", "int", "Unidades por caixa."),
            ("cx_por_pallet", "int", "Caixas por pallet. Deixe 0/vazio se não paletiza (sem arredondamento de pallet)."),
            ("estoque_un", "int", "Estoque atual em UNIDADES (soma dos locais que contam como cobertura)."),
            ("em_producao", "int", "Unidades já em produção/ordem aberta."),
            ("meta_mensal", "int", "Meta de vendas do mês em UNIDADES (venda diária = meta ÷ 30)."),
            ("reservado_cx", "int", "Opcional: caixas reservadas (não contam no estoque, se você marcar 'descontar reservado')."),
            ("prioridade", "int", "Opcional: ordem de produção quando o material é compartilhado (1 = primeiro). Vazio = menor cobertura primeiro."),
        ],
        "exemplo": ["EXEMPLO-001", "EXEMPLO — substituir", "Plenno", 12, 92, 0, 0, 0, 0, ""],
        "validacoes": {"fabrica": LOCAIS},
    },
    {
        "nome": "Estoque_Insumos", "cor": VERDE_ESCURO, "quem": "Você (pede ao terceiro)", "freq": "Todo mês, com a data",
        "texto": "Saldo de matéria-prima e embalagem de CADA local (Plenno, LiveMax, Bioz). Uma linha por insumo e local. "
                 "Sempre informe a data do saldo: o PDF e o alerta de saldo velho usam essa data.",
        "colunas": [
            ("comp_sku", "txt", "Código do insumo (igual à BOM e ao Cadastro)."),
            ("descricao", "desc", "Preenche sozinha a partir da aba Cadastro (fórmula) — não digite."),
            ("estoque", "num", "Saldo na MESMA unidade da BOM (UND, KG, L…)."),
            ("local", "txt", "Onde o material está: Plenno, LiveMax ou Bioz."),
            ("data_saldo", "data", "Data em que o terceiro informou o saldo (DD/MM/AAAA)."),
        ],
        "exemplo": ["EXEMPLO-INS", "", 0, "Plenno", ""],
        "validacoes": {"local": LOCAIS},
        "formula_desc": True,
    },
    {
        "nome": "Pedidos_Criados", "cor": VERDE_ESCURO, "quem": "Você", "freq": "A cada pedido",
        "texto": "Pedidos de compra de matéria-prima/embalagem que você criou, com a DATA DE ENTREGA e para qual terceiro vai. "
                 "A quantidade é a que você decidiu (pode usar a sugestão do MRP e arredondar). O sistema usa a data para dizer "
                 "quando cada item poderá ser produzido.",
        "colunas": [
            ("oc", "txt", "Nº do pedido de compra."),
            ("comp_sku", "txt", "Código do insumo (igual à BOM)."),
            ("descricao", "desc", "Preenche sozinha a partir da aba Cadastro (fórmula) — não digite."),
            ("qtd", "num", "Quantidade do pedido, na unidade da BOM (a que você arredondou/alterou)."),
            ("data_entrega", "data", "Data prevista de chegada no terceiro (DD/MM/AAAA)."),
            ("entrega_em", "txt", "Para onde o material vai: Plenno, LiveMax ou Bioz."),
            ("fornecedor", "txt", "Fornecedor do insumo."),
        ],
        "exemplo": ["EXEMPLO-OC", "EXEMPLO-INS", "", 0, "", "Plenno", "EXEMPLO — substituir"],
        "validacoes": {"entrega_em": LOCAIS},
        "formula_desc": True,
    },
    {
        "nome": "Cadastro", "cor": VERDE_ESCURO, "quem": "Você", "freq": "Quando mudar",
        "texto": "Um insumo por linha: descrição, unidade, preço de custo e fornecedor principal. O preço é o que dá o valor do estoque e da compra.",
        "colunas": [
            ("comp_sku", "txt", "Código do insumo."),
            ("descricao", "desc", "Descrição."),
            ("unidade", "txt", "Unidade (UND, KG, L, CX, MI…)."),
            ("preco", "brl", "Preço de custo por unidade."),
            ("fornecedor", "txt", "Fornecedor principal."),
        ],
        "exemplo": ["EXEMPLO-INS", "EXEMPLO — substituir", "UND", 0, "EXEMPLO — substituir"],
    },
    {
        "nome": "Lead_Time", "cor": AMARELO, "quem": "Você", "freq": "Quando mudar",
        "texto": "Prazo de fabricação + entrega de cada fornecedor, e lote mínimo/múltiplo de compra (0 = sem). "
                 "Vazio = o sistema usa o lead time padrão.",
        "colunas": [
            ("Fornecedor", "txt", "Nome igual ao do Cadastro."),
            ("Lead Time (dias)", "int", "Dias entre pedir e receber."),
            ("Lote Mínimo", "num", "Menor quantidade que o fornecedor vende (0 = sem)."),
            ("Múltiplo de Compra", "num", "A compra tem que ser múltiplo disto (0 = sem)."),
        ],
        "exemplo": ["EXEMPLO — substituir", 30, 0, 0],
    },
    {
        "nome": "Capacidade", "cor": AMARELO, "quem": "Você", "freq": "Quando mudar",
        "texto": "Máximo de posições de pallet de cada fábrica/terceiro. O sistema avisa quando a produção planejada estoura.",
        "colunas": [
            ("Fábrica", "txt", "Plenno, LiveMax ou Bioz (igual à aba SKUs)."),
            ("Capacidade (posições de pallet)", "int", "Posições de pallet disponíveis."),
        ],
        "exemplo": ["EXEMPLO — substituir", 0],
    },
    {
        "nome": "Metas", "cor": AMARELO, "quem": "Você", "freq": "Todo mês",
        "texto": "Opcional. Meta de vendas por mês, para simular com a meta de outro mês sem mexer na aba SKUs "
                 "(ex.: em novembro, planejar dezembro). PRODUTO igual à descrição da aba SKUs.",
        "colunas": [("PRODUTO", "desc", "Igual à coluna descricao da aba SKUs.")]
                   + [(m, "int", "Meta do mês em unidades.") for m in
                      ["JANEIRO", "FEVEREIRO", "MARCO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO",
                       "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]],
        "exemplo": ["EXEMPLO — substituir"] + [0] * 12,
    },
    {
        "nome": "BOM", "cor": LARANJA, "quem": "Tiny (estruturas)", "freq": "Quando mudar a estrutura",
        "texto": "ÚNICA aba que vem do Tiny: a estrutura de cada produto (quanto de cada insumo vai em 1 unidade). "
                 "Por enquanto cole aqui o que o Tiny exporta; depois pode vir por API. Não digite à mão.",
        "colunas": [
            ("kit_sku", "txt", "Código do produto acabado/kit (igual à coluna sku da aba SKUs)."),
            ("comp_sku", "txt", "Código do componente (insumo)."),
            ("comp_id", "txt", "ID do componente no Tiny (ajuda a casar pedidos antigos)."),
            ("comp_desc", "desc", "Descrição do componente."),
            ("qtd", "num", "Quantidade do componente por 1 unidade do produto."),
            ("un", "txt", "Unidade do componente."),
        ],
        "exemplo": ["EXEMPLO-001", "EXEMPLO-INS", "", "EXEMPLO — substituir", 1, "UND"],
    },
]


def _cab(cel, cor=VERDE_ESCURO):
    cel.font = Font(name="Calibri", bold=True, color="FFFFFF")
    cel.fill = PatternFill("solid", fgColor=cor)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border = BORDA


def _leia_me(wb: Workbook, preenchido: bool) -> None:
    ws = wb.active
    ws.title = "LEIA-ME"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Bioz Green — Planilha de entrada do Planejador MRP/PCP"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=VERDE_ESCURO)
    ws["A2"] = "Bioz Green — Pelo Planeta"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color=VERDE_CLARO)
    ws.merge_cells("A1:E1")
    ws.merge_cells("A2:E2")
    nota = ("Esta planilha já vem com os seus dados atuais (Plenno, Dados ponto Reposição). Atualize o que mudou no mês."
            if preenchido else
            "Substitua a linha de EXEMPLO (fundo creme, letra cinza) pelos seus dados. Linhas de EXEMPLO são ignoradas pelo sistema.")
    passos = [
        "COMO USAR",
        "1. " + nota,
        "2. Não renomeie as abas nem os cabeçalhos da linha 1 (o sistema reconhece as abas pelas colunas).",
        "3. Salve o arquivo na pasta do mês e envie no app (ou deixe no Google Drive).",
        "4. O app lê tudo, calcula e mostra o que produzir, o que comprar e o que cada terceiro já pode produzir.",
        "",
        "CORES DAS ABAS",
        "Verde escuro = você preenche todo mês · Amarelo = você preenche quando mudar · Laranja = vem do Tiny (estruturas).",
        "Células em creme = campos de digitação. Passe o mouse sobre o título de cada coluna para ver a explicação.",
        "",
        "REGRAS DE PREENCHIMENTO",
        "• Datas em DD/MM/AAAA. Decimais com vírgula ou ponto (ex.: 12,5).",
        "• Estoque dos acabados em UNIDADES; estoque e pedidos dos insumos na MESMA unidade da BOM.",
        "• Local/Fábrica sempre um destes: Plenno, LiveMax, Bioz (a lista suspensa já oferece).",
        "• O código do insumo (comp_sku) tem que ser igual nas abas BOM, Estoque_Insumos, Pedidos_Criados e Cadastro.",
        "• Saldo de material: sempre com a data (coluna data_saldo). O alerta de 'saldo velho' usa essa data.",
    ]
    r = 4
    for t in passos:
        c = ws.cell(row=r, column=1, value=t)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if t in ("COMO USAR", "CORES DAS ABAS", "REGRAS DE PREENCHIMENTO"):
            c.font = Font(name="Calibri", bold=True, size=12, color=VERDE_ESCURO)
        else:
            c.font = Font(name="Calibri", size=11, color="3D3D3D")
            if len(t) > 110:
                ws.row_dimensions[r].height = 32
        r += 1
    r += 1
    for j, h in enumerate(["Aba", "Quem preenche", "Quando", "O que colocar"], start=1):
        _cab(ws.cell(row=r, column=j, value=h))
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    for a in ABAS:
        r += 1
        for j, v in enumerate([a["nome"], a["quem"], a["freq"], a["texto"]], start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDA
            c.font = Font(name="Calibri", size=10, bold=(j == 1), color="3D3D3D")
            if j == 1:
                c.fill = PatternFill("solid", fgColor=a["cor"])
                c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF" if a["cor"] != AMARELO else "3D3D3D")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 58
    for col, larg in zip("ABCDE", (22, 24, 24, 62, 30)):
        ws.column_dimensions[col].width = larg
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.cell(row=r + 2, column=1, value="Contato: compras@biozgreen.eco.br").font = Font(name="Calibri", size=10, color=CINZA)


def _frames(dados) -> dict:
    """Converte o objeto Dados nas tabelas de cada aba (colunas na ordem do modelo)."""
    f: dict = {}
    s = dados.skus
    f["SKUs"] = s[["sku", "descricao", "fabrica", "un_por_cx", "cx_por_pallet", "estoque_un", "em_producao",
                   "meta_mensal", "reservado_cx", "prioridade"]].copy()
    if dados.estoque is not None:
        e = dados.estoque.copy()
        e["descricao"] = ""
        f["Estoque_Insumos"] = e[["comp_sku", "descricao", "estoque", "local", "data_saldo"]]
    if dados.pedidos_criados is not None and len(dados.pedidos_criados):
        p = dados.pedidos_criados.copy()
        p["descricao"] = ""
        f["Pedidos_Criados"] = p[["oc", "comp_sku", "descricao", "qtd", "data_entrega", "entrega_em", "fornecedor"]]
    if dados.cadastro is not None:
        f["Cadastro"] = dados.cadastro[["comp_sku", "descricao", "unidade", "preco", "fornecedor"]].copy()
    lt = dados.lead_time
    f["Lead_Time"] = pd.DataFrame({"Fornecedor": lt["fornecedor"], "Lead Time (dias)": lt["lead_time_dias"],
                                   "Lote Mínimo": lt["lote_minimo"], "Múltiplo de Compra": lt["multiplo"]})
    cap = dados.capacidade
    f["Capacidade"] = pd.DataFrame({"Fábrica": cap["fabrica"], "Capacidade (posições de pallet)": cap["capacidade_posicoes"]})
    if dados.metas is not None:
        from engine import norm
        m = dados.metas.copy()
        nome = {norm(d): d for d in s["descricao"]}
        m.insert(0, "PRODUTO", [nome.get(k, k) for k in m.index])
        f["Metas"] = m.reset_index(drop=True)
    if dados.bom is not None:
        b = dados.bom.copy()
        f["BOM"] = b[["kit_sku", "comp_sku", "comp_id", "comp_desc", "qtd", "un"]]
    return f


def _escrever_aba(wb: Workbook, cfg: dict, df: pd.DataFrame | None) -> None:
    ws = wb.create_sheet(cfg["nome"])
    ws.sheet_properties.tabColor = cfg["cor"]
    cols = cfg["colunas"]
    for j, (nome, tipo, ajuda) in enumerate(cols, start=1):
        cel = ws.cell(row=1, column=j, value=nome)
        _cab(cel, LARANJA if cfg["cor"] == LARANJA else VERDE_ESCURO)
        cel.comment = Comment(ajuda, "Bioz Green")
        fmt, larg = FMT[tipo]
        ws.column_dimensions[get_column_letter(j)].width = max(larg, min(len(nome) + 4, 34))
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    n_ent = LINHAS_ENTRADA if cfg["nome"] not in ("BOM",) else 0
    if df is not None:
        linhas = df.where(pd.notna(df), None).values.tolist()
        n_ent = max(n_ent, len(linhas) + 20) if cfg["nome"] != "BOM" else 0
    else:
        linhas = [cfg["exemplo"]]
    total = max(len(linhas) + (20 if cfg["nome"] == "BOM" else 0), n_ent)
    for i in range(total):
        row = i + 2
        v = linhas[i] if i < len(linhas) else [None] * len(cols)
        for j, ((nome, tipo, _), val) in enumerate(zip(cols, v), start=1):
            if isinstance(val, (np.integer,)):
                val = int(val)
            elif isinstance(val, (np.floating,)):
                val = float(val)
            elif isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            if val == "" or (isinstance(val, float) and np.isnan(val)):
                val = None
            cel = ws.cell(row=row, column=j, value=val)
            cel.fill = PatternFill("solid", fgColor=CREME)
            cel.border = BORDA
            fmt = FMT[tipo][0]
            if fmt:
                cel.number_format = fmt
            if df is None and i == 0:
                cel.font = Font(name="Calibri", italic=True, color=CINZA)
            if cfg.get("formula_desc") and nome == "descricao":
                cel.value = f'=IF(A{row}="","",IFERROR(VLOOKUP(A{row},Cadastro!$A:$B,2,FALSE),""))' if cfg["nome"] == "Estoque_Insumos" \
                    else f'=IF(B{row}="","",IFERROR(VLOOKUP(B{row},Cadastro!$A:$B,2,FALSE),""))'
                cel.font = Font(name="Calibri", color=CINZA)
                cel.fill = PatternFill("solid", fgColor="FFFFFF")
    for nome_col, lista in (cfg.get("validacoes") or {}).items():
        j = [c[0] for c in cols].index(nome_col) + 1
        dv = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{total + 1}")


def gerar_modelo_xlsx(dados=None) -> bytes:
    wb = Workbook()
    _leia_me(wb, dados is not None)
    frames = _frames(dados) if dados is not None else {}
    for cfg in ABAS:
        df = frames.get(cfg["nome"]) if dados is not None else None
        if dados is not None and df is None:
            df = pd.DataFrame(columns=[c[0] for c in cfg["colunas"]])
        _escrever_aba(wb, cfg, df)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    with open("modelo_entrada_bioz.xlsx", "wb") as f:
        f.write(gerar_modelo_xlsx())
    print("modelo_entrada_bioz.xlsx gerado")
