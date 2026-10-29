"""Exporta o resultado da simulação e do MRP para Excel (.xlsx) na identidade Bioz Green."""
from __future__ import annotations

import io
from typing import Dict, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERDE_ESCURO = "0B592F"
VERDE_CLARO = "8CA70A"
CREME = "F0E7D7"
_b = Side(style="thin", color="D9D9D9")
_borda = Border(top=_b, bottom=_b, left=_b, right=_b)

ROTULOS = {
    "sku": "SKU", "descricao": "Descrição", "fabrica": "Fábrica", "un_por_cx": "Unid./caixa",
    "cx_por_pallet": "Caixas/pallet", "un_por_pallet": "Unid./pallet", "estoque_un": "Estoque (un.)",
    "estoque_cx": "Estoque (cx)", "em_producao": "Em produção (un.)", "meta_mensal": "Meta mensal",
    "venda_diaria": "Venda diária", "cob_atual_dias": "Cobertura atual (dias)",
    "cob_pos_prod_dias": "Cobertura pós produção (dias)", "faixa": "Faixa", "dias_alvo": "Dias alvo",
    "nec_un": "Necessário (un.)", "nec_cx": "Necessário (cx)", "cx_mult_pallet": "Caixas (múltiplo pallet)",
    "un_a_produzir": "Unidades a produzir", "pallets": "Pallets", "cob_final_dias": "Cobertura final (dias)",
    "pos_estoque": "Posições estoque", "pos_em_producao": "Posições em produção",
    "pos_producao": "Posições produção simulada", "ajuste_manual": "Ajuste manual", "alerta_dados": "Alertas",
    "comp_sku": "Insumo (SKU)", "unidade": "Unidade", "fornecedor": "Fornecedor",
    "consumo_bruto": "Consumo bruto", "n_skus": "Nº SKUs", "estoque": "Estoque insumo",
    "transito_total": "Em trânsito (total)", "transito_considerado": "Em trânsito (considerado)",
    "transito_fora_prazo": "Em trânsito (chega após a data)", "chegada_prevista": "Chegada prevista",
    "necessidade_liquida": "Necessidade líquida", "qtd_comprar": "Qtd a comprar",
    "qtd_sugerida": "Qtd sugerida (lote/múltiplo)", "preco": "Preço unit.", "valor_estimado": "Valor estimado (R$)",
    "lead_time_dias": "Lead time (dias)", "lead_time_padrao": "Lead time é o padrão?",
    "data_necessaria": "Data necessária", "data_limite_pedido": "Data limite do pedido",
    "dias_ate_limite": "Dias até o limite", "status": "Status",
    "sem_estoque_cadastrado": "Sem estoque cadastrado", "sem_fornecedor": "Sem fornecedor", "sem_preco": "Sem preço",
    "dias_alvo_c": "Dias alvo", "skus_a_produzir": "SKUs a produzir", "unidades": "Unidades", "caixas": "Caixas",
    "posicoes_pallet": "Posições de pallet", "insumos_a_comprar": "Insumos a comprar",
    "insumos_atrasados": "Insumos atrasados", "valor_compras": "Valor das compras (R$)",
    "capacidade_posicoes": "Capacidade (posições)", "ocupacao": "Ocupação (posições)", "saldo": "Saldo",
    "uso_pct": "Uso da capacidade",
}

FORMATOS = {
    "preco": '#,##0.00', "valor_estimado": '"R$" #,##0.00', "valor_compras": '"R$" #,##0.00',
    "uso_pct": "0%", "data_necessaria": "DD/MM/YYYY", "data_limite_pedido": "DD/MM/YYYY",
    "chegada_prevista": "DD/MM/YYYY",
}


def _escrever(ws, df: pd.DataFrame, titulo: str, subtitulo: str = "", plano: bool = False):
    """plano=True: cabeçalho na linha 1, sem títulos — formato que o app consegue ler de volta."""
    if plano:
        topo = 1
    else:
        ws["A1"] = titulo
        ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=VERDE_ESCURO)
        ws["A2"] = "Bioz Green — Pelo Planeta" + (f"  |  {subtitulo}" if subtitulo else "")
        ws["A2"].font = Font(name="Calibri", size=10, bold=True, color=VERDE_CLARO)
        topo = 4
    cols = list(df.columns)
    for j, c in enumerate(cols, start=1):
        cel = ws.cell(row=topo, column=j, value=ROTULOS.get(c, c))
        cel.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=VERDE_ESCURO)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = _borda
    for i, row in enumerate(df.itertuples(index=False), start=topo + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, pd.Timestamp):
                v = None if pd.isna(v) else v.to_pydatetime()
            elif isinstance(v, float) and pd.isna(v):
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            cel = ws.cell(row=i, column=j, value=v)
            cel.border = _borda
            fmt = FORMATOS.get(cols[j - 1])
            if fmt:
                cel.number_format = fmt
            elif isinstance(v, float):
                cel.number_format = "#,##0.00"
            elif isinstance(v, int) and not isinstance(v, bool):
                cel.number_format = "#,##0"
    for j, c in enumerate(cols, start=1):
        larg = max([len(str(ROTULOS.get(c, c)))] + [len(str(x)) for x in df[c].head(60).tolist()]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(larg, 11), 48)
    ws.row_dimensions[topo].height = 32
    ws.freeze_panes = ws.cell(row=topo + 1, column=3)
    ws.auto_filter.ref = f"A{topo}:{get_column_letter(len(cols))}{topo + max(len(df), 1)}"


def exportar_excel(
    sim: pd.DataFrame,
    mrp: pd.DataFrame,
    cenarios: Optional[pd.DataFrame] = None,
    ocupacao: Optional[pd.DataFrame] = None,
    avisos: Optional[list] = None,
    lead_time: Optional[pd.DataFrame] = None,
    capacidade: Optional[pd.DataFrame] = None,
    resumo: Optional[Dict[str, str]] = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Planejador MRP/PCP — Bioz Green"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=VERDE_ESCURO)
    ws["A2"] = "Bioz Green — Pelo Planeta"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color=VERDE_CLARO)
    r = 4
    for k, v in (resumo or {}).items():
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=v)
        c.fill = PatternFill("solid", fgColor=CREME)
        r += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 40
    ws.cell(row=r + 1, column=1, value="Contato: compras@biozgreen.eco.br").font = Font(color="595959", size=9)

    _escrever(wb.create_sheet("Simulação PCP"), sim.drop(columns=["ajuste_manual"], errors="ignore"),
              "Simulação de reposição de produtos acabados")
    _escrever(wb.create_sheet("MRP Compras"), mrp, "Necessidade de compra de insumos (MRP)")
    if ocupacao is not None and len(ocupacao):
        _escrever(wb.create_sheet("Pallets por fábrica"), ocupacao, "Ocupação de posições de pallet")
    if cenarios is not None and len(cenarios):
        _escrever(wb.create_sheet("Cenários"), cenarios, "Comparativo de cenários (dias de cobertura alvo)")
    if lead_time is not None and len(lead_time):
        lt = lead_time.rename(columns={"fornecedor": "Fornecedor", "lead_time_dias": "Lead Time (dias)",
                                       "lote_minimo": "Lote Mínimo", "multiplo": "Múltiplo de Compra"})
        _escrever(wb.create_sheet("Lead_Time"), lt, "Lead time por fornecedor", plano=True)
    if capacidade is not None and len(capacidade):
        cp = capacidade.rename(columns={"fabrica": "Fábrica", "capacidade_posicoes": "Capacidade (posições de pallet)"})
        _escrever(wb.create_sheet("Capacidade"), cp, "Capacidade de pallets por fábrica", plano=True)
    if avisos:
        av = pd.DataFrame(avisos, columns=["Nível", "Mensagem"])
        _escrever(wb.create_sheet("Qualidade dos dados"), av, "Pontos de atenção nos dados de entrada")
        wb["Qualidade dos dados"].column_dimensions["B"].width = 140

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
