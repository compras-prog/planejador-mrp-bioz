"""
Planejador MRP / PCP — Bioz Green  (Streamlit)

Executar localmente:   streamlit run app.py
"""
from __future__ import annotations

import hashlib
import hmac
import io
import os
import re
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import altair as alt
import numpy as np
import pandas as pd
import streamlit as st

import theme
from engine import (BASES_OCUPACAO, FAIXAS_ORDEM, calcular_mrp, comparar_cenarios, consumo_detalhado,
                    cronograma_terceiro, faixa_cobertura, ocupacao_armazem, simular_reposicao)
from export import exportar_excel
from loaders import Planilhas, aplicar_meta, carregar, descobrir
from modelo import gerar_modelo_xlsx
from relatorio_pdf import gerar_pdf_programacao

ASSETS = Path(__file__).parent / "assets"
LOGO = ASSETS / "logo-bioz-green-small.png"
ROTULO_PAPEL = {
    "skus": "SKUs / estoque / meta", "dados_sku": "Caixas e pallets por SKU", "bom": "Estrutura (BOM)",
    "estoque_insumos": "Estoque de insumos", "pedidos": "Pedidos em trânsito", "cadastro": "Cadastro de componentes",
    "metas": "Meta de vendas por mês", "lead_time": "Lead time por fornecedor", "capacidade": "Capacidade de pallets",
}

st.set_page_config(page_title="Planejador MRP/PCP — Bioz Green", page_icon="🌿", layout="wide",
                   initial_sidebar_state="expanded")
st.markdown(theme.CSS, unsafe_allow_html=True)


# --------------------------------------------------------------------------- #
# Acesso
# --------------------------------------------------------------------------- #
def _senha_configurada() -> str | None:
    try:
        s = st.secrets.get("APP_PASSWORD", None)
    except Exception:
        s = None
    return s or os.environ.get("APP_PASSWORD")


def _login():
    senha = _senha_configurada()
    if not senha or st.session_state.get("auth_ok"):
        return
    st.markdown('<div class="bz-topo"></div><div class="bz-fio"></div>', unsafe_allow_html=True)
    st.title("Planejador MRP / PCP")
    st.caption("Bioz Green — Pelo Planeta")
    with st.form("login"):
        digitada = st.text_input("Senha de acesso", type="password")
        if st.form_submit_button("Entrar"):
            if hmac.compare_digest(digitada.encode(), senha.encode()):
                st.session_state["auth_ok"] = True
                st.rerun()
            else:
                st.error("Senha incorreta.")
    st.stop()


# --------------------------------------------------------------------------- #
# Dados
# --------------------------------------------------------------------------- #
@st.cache_data(show_spinner="Lendo planilhas…", max_entries=4)
def _carregar(arquivos: tuple, escolhas: tuple):
    pl = Planilhas(dict(arquivos))
    cand = descobrir(pl)
    esc = {p: (a, s) for p, a, s in escolhas if (a, s) in cand.get(p, [])}
    return carregar(pl, esc), cand


def _arquivos_da_fonte() -> dict:
    """Sidebar: escolhe a fonte e devolve {nome: bytes}."""
    opcoes = ["Enviar arquivos", "Google Drive"] + (["Pasta local"] if os.environ.get("MRP_FILES") else [])
    fonte = st.radio("Fonte dos dados", opcoes, label_visibility="collapsed")
    arquivos: dict = {}
    if fonte == "Enviar arquivos":
        ups = st.file_uploader("Planilhas (.xlsx)", type=["xlsx"], accept_multiple_files=True,
                               help="Pode enviar 'Dados ponto Reposição', 'Demanda Plenno' e/ou o modelo padronizado. "
                                    "As abas são reconhecidas pelas colunas.")
        arquivos = {u.name: u.getvalue() for u in ups or []}
    elif fonte == "Google Drive":
        import drive  # import tardio: só precisa das libs do Google se usar esta opção

        ids = st.text_area("IDs ou links (um por linha)", placeholder="https://docs.google.com/spreadsheets/d/…",
                           height=90)
        if st.button("Carregar do Drive"):
            try:
                try:
                    info = dict(st.secrets["gcp_service_account"])
                except Exception:
                    info = None
                novos = {}
                for linha in [x for x in ids.splitlines() if x.strip()]:
                    nome, b = drive.baixar_planilha(drive.extrair_id(linha), info)
                    novos[nome] = b
                st.session_state["drive_arquivos"] = novos
            except Exception as e:  # noqa: BLE001
                st.error(f"Não consegui ler do Drive: {e}")
        arquivos = st.session_state.get("drive_arquivos", {})
        if arquivos:
            st.caption("Carregado: " + ", ".join(arquivos))
    else:
        for caminho in os.environ["MRP_FILES"].split(";"):
            p = Path(caminho.strip())
            if p.exists():
                arquivos[p.name] = p.read_bytes()
        st.caption("Pasta local: " + ", ".join(arquivos))
    st.download_button("Baixar modelo padronizado", gerar_modelo_xlsx(), "modelo_entrada_bioz.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
    return arquivos


# --------------------------------------------------------------------------- #
# Gráficos
# --------------------------------------------------------------------------- #
def grafico_faixas(faixas: pd.Series, titulo: str):
    cont = faixas.value_counts().reindex(FAIXAS_ORDEM, fill_value=0).rename_axis("faixa").reset_index(name="SKUs")
    cont = cont[(cont["SKUs"] > 0) | (cont["faixa"] != "Sem meta")]
    ch = alt.Chart(cont, title=titulo).mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4).encode(
        x=alt.X("faixa:N", sort=FAIXAS_ORDEM, title=None, axis=alt.Axis(labelAngle=-25, labelLimit=170, labelOverlap=False)),
        y=alt.Y("SKUs:Q", title="SKUs", axis=alt.Axis(tickMinStep=1)),
        color=alt.Color("faixa:N", scale=alt.Scale(domain=list(theme.COR_FAIXA), range=list(theme.COR_FAIXA.values())),
                        legend=None),
        tooltip=["faixa", "SKUs"],
    ).properties(height=260)
    st.altair_chart(ch, width="stretch")


def grafico_barra(df: pd.DataFrame, x: str, y: str, titulo: str, cor: str, fmt: str = ",.0f"):
    ch = alt.Chart(df, title=titulo).mark_bar(color=cor, cornerRadiusTopLeft=4, cornerRadiusTopRight=4).encode(
        x=alt.X(f"{x}:O", title="Dias de cobertura alvo", axis=alt.Axis(labelAngle=0)),
        y=alt.Y(f"{y}:Q", title=None, axis=alt.Axis(format=fmt)),
        tooltip=[alt.Tooltip(f"{x}:O", title="Dias"), alt.Tooltip(f"{y}:Q", format=fmt)],
    ).properties(height=220)
    st.altair_chart(ch, width="stretch")


def grafico_ocupacao(oc: pd.DataFrame):
    base = alt.Chart(oc).encode(x=alt.X("fabrica:N", title=None, axis=alt.Axis(labelAngle=0)))
    barras = base.mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4, size=60).encode(
        y=alt.Y("ocupacao:Q", title="Posições de pallet"),
        color=alt.condition(alt.datum.status == "ESTOURA CAPACIDADE", alt.value(theme.LARANJA),
                            alt.value(theme.VERDE_ESCURO)),
        tooltip=["fabrica", "ocupacao", "capacidade_posicoes", "status"])
    cap = base.mark_tick(color=theme.AMARELO, thickness=4, size=90).encode(y="capacidade_posicoes:Q")
    st.altair_chart((barras + cap).properties(height=280, title="Ocupação x capacidade (traço amarelo = capacidade)"),
                    width="stretch")


# --------------------------------------------------------------------------- #
# Pedidos de compra criados (helpers)
# --------------------------------------------------------------------------- #
def _po_vazio() -> pd.DataFrame:
    return pd.DataFrame({"oc": pd.Series(dtype="object"), "insumo": pd.Series(dtype="object"),
                         "qtd": pd.Series(dtype="float"), "sugerido_mrp": pd.Series(dtype="float"),
                         "data_entrega": pd.Series(dtype="datetime64[ns]"),
                         "entrega_em": pd.Series(dtype="object"), "fornecedor": pd.Series(dtype="object")})


def _normalizar_po(df: pd.DataFrame, opcoes_ins: dict) -> pd.DataFrame:
    """Aceita o Excel exportado pelo app (ou colunas parecidas) e devolve o formato do editor."""
    from loaders import achar_col, to_date, to_num
    c = {"oc": achar_col(df, "oc", "pedido", "ordem de compra", "numero"), "insumo": achar_col(df, "insumo", "comp_sku", "sku"),
         "qtd": achar_col(df, "quantidade", "qtd"), "data": achar_col(df, "data de entrega", "data_entrega", "data prevista", "prevista"),
         "local": achar_col(df, "entrega em", "entrega_em", "local"), "forn": achar_col(df, "fornecedor"),
         "sug": achar_col(df, "sugerido pelo mrp", "sugerido_mrp")}
    if not c["insumo"] or not c["qtd"]:
        raise ValueError("preciso das colunas de Insumo e Quantidade")
    cod = df[c["insumo"]].astype(str).str.split(" — ").str[0].str.strip()
    out = _po_vazio()
    out = pd.DataFrame({
        "oc": df[c["oc"]].astype(str).replace("nan", "") if c["oc"] else "",
        "insumo": cod.map(lambda k: opcoes_ins.get(k, k)),
        "qtd": to_num(df[c["qtd"]]),
        "sugerido_mrp": to_num(df[c["sug"]]) if c["sug"] else np.nan,
        "data_entrega": pd.to_datetime(to_date(df[c["data"]])) if c["data"] else pd.NaT,
        "entrega_em": df[c["local"]].fillna("").astype(str) if c["local"] else "",
        "fornecedor": df[c["forn"]].fillna("").astype(str) if c["forn"] else "",
    })
    return out


def _po_para_calculo(po_ed: pd.DataFrame) -> pd.DataFrame:
    """Linhas válidas → comp_sku, qtd, data, entrega_em, fornecedor, oc."""
    if po_ed is None or not len(po_ed):
        return pd.DataFrame(columns=["comp_sku", "qtd", "data", "entrega_em", "fornecedor", "oc"])
    d = po_ed.dropna(subset=["insumo"]).copy()
    d = d[pd.to_numeric(d["qtd"], errors="coerce").fillna(0) > 0]
    return pd.DataFrame({
        "comp_sku": d["insumo"].astype(str).str.split(" — ").str[0].str.strip(),
        "qtd": pd.to_numeric(d["qtd"], errors="coerce"), "data": pd.to_datetime(d["data_entrega"], errors="coerce"),
        "entrega_em": d["entrega_em"].fillna("") if "entrega_em" in d else "",
        "fornecedor": d["fornecedor"].fillna("") if "fornecedor" in d else "", "oc": d["oc"].fillna("") if "oc" in d else "",
    }).reset_index(drop=True)


def _po_para_excel(po_ed: pd.DataFrame, buf) -> None:
    d = po_ed.copy().rename(columns={"oc": "OC", "insumo": "Insumo", "qtd": "Quantidade", "sugerido_mrp": "Sugerido pelo MRP", "data_entrega": "Data de entrega",
                                     "entrega_em": "Entrega em", "fornecedor": "Fornecedor"})
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
        d.to_excel(w, index=False, sheet_name="Pedidos_Criados")
        ws = w.sheets["Pedidos_Criados"]
        from openpyxl.styles import Font, PatternFill
        for cel in ws[1]:
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="0B592F")
        for col, larg in zip("ABCDEFG", (16, 60, 14, 18, 16, 14, 34)):
            ws.column_dimensions[col].width = larg


def _montar_sugestoes(mrp_base: pd.DataFrame, sim_compra: pd.DataFrame, bom, fabricas) -> pd.DataFrame:
    """Insumos a comprar segundo o MRP (sem pedidos criados): qtd sugerida arredondada p/ cima em 2 casas + terceiro que consome."""
    cols = ["comp_sku", "qtd", "fornecedor", "entrega_em"]
    if mrp_base is None or not len(mrp_base) or bom is None:
        return pd.DataFrame(columns=cols)
    m = mrp_base[mrp_base["qtd_comprar"] > 0]
    if m.empty:
        return pd.DataFrame(columns=cols)
    prod = sim_compra[sim_compra["un_a_produzir"].fillna(0) > 0][["sku", "fabrica"]]
    if fabricas is not None:
        prod = prod[prod["fabrica"].isin(list(fabricas))]
    uso = bom.merge(prod, left_on="kit_sku", right_on="sku").groupby("comp_sku")["fabrica"].agg(lambda s: sorted(set(s)))
    ent = m["comp_sku"].map(lambda k: uso.get(k, [""])[0] if len(uso.get(k, [])) == 1 else "")
    return pd.DataFrame({"comp_sku": m["comp_sku"].values, "qtd": np.ceil(m["qtd_sugerida"].values * 100 - 1e-9) / 100,
                         "fornecedor": m["fornecedor"].values, "entrega_em": ent.values}).reset_index(drop=True)


def _somar_pedidos(pedidos, po_df: pd.DataFrame) -> pd.DataFrame:
    """Pedidos do ERP + pedidos criados aqui (viram 'componente' com a data de entrega informada)."""
    base = pedidos if pedidos is not None else pd.DataFrame(columns=["comp_sku", "tipo", "qtd", "prevista"])
    if po_df is None or not len(po_df):
        return base
    novos = pd.DataFrame({"oc": po_df["oc"], "fornecedor": po_df["fornecedor"], "item": po_df["comp_sku"],
                          "qtd": po_df["qtd"], "prevista": po_df["data"], "comp_sku": po_df["comp_sku"], "tipo": "componente"})
    return pd.concat([base, novos], ignore_index=True)


def _data_do_saldo(texto_aba: str, hoje):
    """Tenta ler dd_mm do nome da aba (ex.: 'Estoque Plenno_02_10')."""
    aba = (texto_aba or "").split(" › ")[-1].split("; ")[0] if "; " in (texto_aba or "") else (texto_aba or "").split(" › ")[-1]
    m = re.search(r"(\d{2})_(\d{2})", aba)
    if m:
        try:
            return datetime(hoje.year, int(m.group(2)), int(m.group(1))).date()
        except ValueError:
            pass
    return hoje


# --------------------------------------------------------------------------- #
# App
# --------------------------------------------------------------------------- #
def main():
    _login()
    hoje = datetime.now(ZoneInfo("America/Sao_Paulo")).date()

    st.markdown('<div class="bz-topo"></div><div class="bz-fio"></div>', unsafe_allow_html=True)
    c_logo, c_tit = st.columns([1, 8], vertical_alignment="center")
    if LOGO.exists():
        c_logo.image(str(LOGO), width=84)
    c_tit.title("Planejador MRP / PCP")
    c_tit.markdown('<div class="bz-sub">Bioz Green — Pelo Planeta · reposição de acabados e compra de insumos</div>',
                   unsafe_allow_html=True)

    # ---------------- Sidebar: dados -----------------
    with st.sidebar:
        st.header("1. Dados")
        arquivos = _arquivos_da_fonte()

    if not arquivos:
        st.info("**Comece enviando as planilhas** na barra lateral (ou o modelo padronizado). "
                "O app reconhece as abas pelas colunas — não precisa renomear nada.")
        with st.expander("O que o sistema calcula?", expanded=True):
            st.markdown(
                "- **PCP:** cobertura atual e pós-produção, necessidade para os dias-alvo (60/75/80…), "
                "caixas → múltiplo de pallet → unidades a produzir, posições de pallet e alerta de capacidade.\n"
                "- **MRP:** consumo bruto pela BOM, necessidade líquida (estoque + pedidos em trânsito), "
                "data limite de pedido pelo lead time do fornecedor e valor estimado.\n"
                "- **Cenários** lado a lado e **exportação em Excel**."
            )
        st.stop()

    escolhas = []
    for papel in ROTULO_PAPEL:
        v = st.session_state.get(f"esc_{papel}")
        if v and " › " in v:
            a, s = v.split(" › ", 1)
            escolhas.append((papel, a, s))
    try:
        dados, cand = _carregar(tuple(sorted(arquivos.items())), tuple(escolhas))
    except ValueError as e:
        st.error(str(e))
        st.stop()
    if dados.skus.empty:
        st.info("A planilha ainda não tem nenhum SKU preenchido (as linhas de EXEMPLO são ignoradas). "
                "Preencha a aba **SKUs** (e as demais) e envie de novo.")
        st.stop()

    sig = hashlib.md5("|".join(f"{k}:{len(v)}" for k, v in sorted(arquivos.items())).encode()).hexdigest()[:8]
    rst = st.session_state.get("reset_n", 0)

    # ---------------- Sidebar: parâmetros -----------------
    with st.sidebar:
        with st.expander("Abas detectadas", expanded=False):
            for papel, rot in ROTULO_PAPEL.items():
                opcoes = [f"{a} › {s}" for a, s in cand.get(papel, [])]
                if len(opcoes) > 1:
                    st.selectbox(rot, opcoes, key=f"esc_{papel}")
                else:
                    st.caption(f"**{rot}:** {opcoes[0] if opcoes else '— não encontrada —'}")

        st.header("2. Parâmetros")
        opc = st.radio("Dias de cobertura alvo", [60, 75, 80, "Outro"], horizontal=True)
        dias = st.number_input("Dias alvo", 7, 365, 90, step=5) if opc == "Outro" else int(opc)
        dias_cen = st.multiselect("Cenários para comparar", [30, 45, 60, 75, 80, 90, 120], default=[60, 75, 80])
        dias_cen = sorted(set(dias_cen) | {dias})
        dias_compra_sel = st.selectbox("Comprar insumos para (dias)", ["Igual à produção", 45, 60, 75, 90, 120],
                                       help="Horizonte do MRP de compras. Ex.: produzir para 40 dias, mas já comprar "
                                            "material para 60.")

        meses = [] if dados.metas is None else list(dados.metas.columns)
        opcoes_meta = ["Meta do mês (planilha de reposição)"] + [f"Meta de {m.title()}" for m in meses]
        sel_meta = st.selectbox("Meta de vendas", opcoes_meta)
        mes_sel = None if sel_meta == opcoes_meta[0] else meses[opcoes_meta.index(sel_meta) - 1]
        ajuste = st.slider("Ajuste na demanda (%)", -50, 100, 0, step=5, help="Cenário: +10 simula demanda 10% maior.")
        mult = 1 + ajuste / 100
        desc_res = st.checkbox("Descontar estoque reservado", value=False)

        fabs = sorted(dados.skus["fabrica"].unique())
        fab_mrp = st.multiselect("Fábricas incluídas no MRP", fabs, default=fabs)
        data_nec = st.date_input("Início planejado da produção", value=hoje + timedelta(days=7), format="DD/MM/YYYY",
                                 help="Data em que os insumos precisam estar na fábrica. "
                                      "Data limite do pedido = esta data − lead time do fornecedor.")
        lt_padrao = int(st.number_input("Lead time padrão (dias)", 0, 365, 30,
                                        help="Usado para fornecedores sem lead time informado."))
        transito_ate = st.checkbox("Contar em trânsito só o que chega até a data acima", value=True)
        base_oc = st.selectbox("Ocupação de pallets considera", list(BASES_OCUPACAO), format_func=BASES_OCUPACAO.get)
        if st.button("Restaurar ajustes manuais", width="stretch"):
            st.session_state["reset_n"] = rst + 1
            st.rerun()

    nomes_abas = ["Visão geral", "Simulação PCP", "Pallets por fábrica", "Programação do terceiro",
                  "Pedidos de compra criados", "MRP Compras", "Consumo por insumo", "Parâmetros", "Qualidade dos dados"]
    T = dict(zip(nomes_abas, st.tabs(nomes_abas)))

    # ---------------- Parâmetros editáveis (renderizados primeiro: alimentam o cálculo) -----------------
    skus_base, sem_meta_mes = aplicar_meta(dados.skus, dados.metas, mes_sel)
    with T["Parâmetros"]:
        st.subheader("Lead time, lote mínimo e múltiplo de compra por fornecedor")
        st.caption("Campo vazio = usa o lead time padrão da barra lateral. As alterações valem para esta sessão — "
                   "baixe o Excel (aba Lead_Time) para guardar e reenviar depois.")
        lead_df = st.data_editor(
            dados.lead_time, key=f"lt_{sig}_{rst}", num_rows="dynamic", hide_index=True, width="stretch",
            column_config={
                "fornecedor": st.column_config.TextColumn("Fornecedor", width="large"),
                "lead_time_dias": st.column_config.NumberColumn("Lead time (dias)", min_value=0, step=1),
                "lote_minimo": st.column_config.NumberColumn("Lote mínimo", min_value=0),
                "multiplo": st.column_config.NumberColumn("Múltiplo de compra", min_value=0),
            })
        lead_df = lead_df[lead_df["fornecedor"].fillna("").str.strip() != ""]

        st.subheader("Capacidade de posições de pallet por fábrica / terceiro")
        cap_df = st.data_editor(
            dados.capacidade, key=f"cap_{sig}_{rst}", num_rows="dynamic", hide_index=True, width="stretch",
            column_config={
                "fabrica": st.column_config.TextColumn("Fábrica (igual ao nome na planilha de SKUs)"),
                "capacidade_posicoes": st.column_config.NumberColumn("Capacidade (posições)", min_value=0, step=1),
            })
        cap_df = cap_df[cap_df["fabrica"].fillna("").str.strip() != ""]

        st.subheader("Premissas por SKU (opcional)")
        st.caption("Edite meta, estoque, em produção ou force uma quantidade de **caixas** a produzir "
                   "(substitui o cálculo, sem arredondar para pallet).")
        ed = skus_base[["sku", "descricao", "fabrica", "meta_mensal", "em_producao", "estoque_un"]].copy()
        ed["ajuste_cx"] = np.nan
        ed_out = st.data_editor(
            ed, key=f"sk_{sig}_{mes_sel}_{rst}", hide_index=True, width="stretch", height=380,
            disabled=["sku", "descricao", "fabrica"],
            column_config={
                "sku": "SKU", "descricao": st.column_config.TextColumn("Descrição", width="large"),
                "fabrica": "Fábrica",
                "meta_mensal": st.column_config.NumberColumn("Meta mensal", min_value=0),
                "em_producao": st.column_config.NumberColumn("Em produção (un.)", min_value=0),
                "estoque_un": st.column_config.NumberColumn("Estoque (un.)"),
                "ajuste_cx": st.column_config.NumberColumn("Forçar caixas", min_value=0, step=1),
            })

    skus_edit = skus_base.set_index("sku")
    eo = ed_out.set_index("sku")
    for col in ("meta_mensal", "em_producao", "estoque_un", "ajuste_cx"):
        skus_edit[col] = eo[col]
    skus_edit = skus_edit.reset_index()

    # ---------------- Cálculo -----------------
    sim = simular_reposicao(skus_edit, dias, mult, desc_res)
    ocup = ocupacao_armazem(sim, cap_df, base_oc)

    tem_bom = dados.bom is not None
    kw = dict(
        bom=dados.bom if tem_bom else pd.DataFrame(columns=["kit_sku", "comp_sku", "comp_id", "comp_desc", "qtd", "un"]),
        estoque=dados.estoque if dados.estoque is not None else pd.DataFrame(columns=["comp_sku", "estoque"]),
        pedidos=_somar_pedidos(dados.pedidos, None),
        cadastro=dados.cadastro if dados.cadastro is not None
        else pd.DataFrame(columns=["comp_sku", "descricao", "unidade", "preco", "fornecedor"]),
        parametros_forn=lead_df, data_necessaria=data_nec, hoje=hoje, lt_padrao=lt_padrao,
        transito_ate_data=transito_ate, fabricas=fab_mrp,
    )
    dias_compra = dias if dias_compra_sel == "Igual à produção" else int(dias_compra_sel)
    sim_compra = sim if dias_compra == dias else simular_reposicao(skus_edit.drop(columns=["ajuste_cx"]), dias_compra,
                                                                    mult, desc_res)
    # MRP "base" = sem os pedidos criados aqui: é a SUGESTÃO que alimenta o cadastro de pedidos
    mrp_base = calcular_mrp(sim_compra, **kw) if tem_bom else pd.DataFrame()
    sugestoes = _montar_sugestoes(mrp_base, sim_compra, dados.bom if tem_bom else None, fab_mrp)

    # ---------------- Pedidos de compra criados (material que vai chegar) — alimenta MRP e programação -----------------
    opcoes_ins = {}
    if dados.bom is not None:
        for r in dados.bom.drop_duplicates("comp_sku").itertuples():
            opcoes_ins[r.comp_sku] = f"{r.comp_sku} — {r.comp_desc}"
    locais_po = [""] + sorted(dados.skus["fabrica"].unique())
    with T["Pedidos de compra criados"]:
        st.subheader("Pedidos de compra criados — material que vai chegar nos terceiros")
        st.caption("Cadastre aqui cada pedido que você emitir, com a **data de entrega** no terceiro. A aba "
                   "*Programação do terceiro* usa essas datas para simular quando cada item pode ser produzido, e o MRP "
                   "passa a contar esse material como \"a caminho\". Salve o Excel na pasta do mês e recarregue depois.")
        pc = dados.pedidos_criados
        if pc is not None and len(pc):                       # aba Pedidos_Criados da planilha de entrada
            h0 = hashlib.md5(pc.to_csv(index=False).encode()).hexdigest()
            if st.session_state.get("po_pc_hash") != h0:
                st.session_state["po_base"] = _normalizar_po(pc, opcoes_ins)
                st.session_state["po_pc_hash"] = h0
                st.session_state.pop("po_last", None)
                st.session_state["po_ver"] = st.session_state.get("po_ver", 0) + 1
                st.info(f"{len(pc)} pedido(s) lido(s) da aba Pedidos_Criados da planilha de entrada.")
        up = st.file_uploader("Carregar pedidos já criados (.xlsx ou .csv)", type=["xlsx", "csv"], key="po_up")
        if up is not None:
            h = hashlib.md5(up.getvalue()).hexdigest()
            if st.session_state.get("po_hash") != h:
                try:
                    st.session_state["po_base"] = _normalizar_po(
                        pd.read_csv(io.BytesIO(up.getvalue()), sep=None, engine="python") if up.name.lower().endswith(".csv")
                        else pd.read_excel(io.BytesIO(up.getvalue())), opcoes_ins)
                    st.session_state["po_hash"] = h
                    st.session_state["po_ver"] = st.session_state.get("po_ver", 0) + 1
                except Exception as e:  # noqa: BLE001
                    st.error(f"Não consegui ler o arquivo: {e}")
        if len(sugestoes):
            st.markdown(f"**Sugestão do MRP:** {len(sugestoes)} insumos a comprar para {dias_compra} dias de produção "
                        f"(já com lote mínimo e múltiplo). Preencha o pedido com a sugestão e **ajuste as quantidades** "
                        f"como quiser — o que valer é o que estiver na coluna *Quantidade*.")
            f1, f2 = st.columns([4, 2])
            forn_sel = f1.selectbox("Fornecedor do pedido", ["(todos)"] + sorted(sugestoes["fornecedor"].unique()),
                                    key="po_forn")
            if f2.button("Preencher com a sugestão do MRP", width="stretch"):
                sg = sugestoes if forn_sel == "(todos)" else sugestoes[sugestoes["fornecedor"] == forn_sel]
                atual = st.session_state.get("po_last", st.session_state.get("po_base", _po_vazio()))
                ja = set(atual["insumo"].dropna().astype(str).str.split(" — ").str[0]) if len(atual) else set()
                novos = sg[~sg["comp_sku"].isin(ja)]
                linhas_novas = pd.DataFrame({
                    "oc": "", "insumo": novos["comp_sku"].map(lambda k: opcoes_ins.get(k, k)),
                    "qtd": novos["qtd"], "sugerido_mrp": novos["qtd"], "data_entrega": pd.NaT,
                    "entrega_em": novos["entrega_em"], "fornecedor": novos["fornecedor"]})
                st.session_state["po_base"] = pd.concat([atual, linhas_novas], ignore_index=True) if len(atual) \
                    else linhas_novas
                st.session_state["po_ver"] = st.session_state.get("po_ver", 0) + 1
                st.rerun()
        else:
            st.caption("O MRP não sugere compras para o cenário atual (ou faltam BOM/estoque).")
        base_po = st.session_state.get("po_base", _po_vazio())
        po_ed = st.data_editor(
            base_po, key=f"po_{st.session_state.get('po_ver', 0)}_{rst}", num_rows="dynamic", hide_index=True, width="stretch",
            column_config={
                "oc": st.column_config.TextColumn("Nº do pedido (OC)"),
                "insumo": st.column_config.SelectboxColumn("Insumo", options=list(opcoes_ins.values()), width="large",
                                                           required=True),
                "qtd": st.column_config.NumberColumn("Quantidade do pedido", min_value=0,
                                                     help="Edite à vontade (arredonde, some, reduza). Unidade = a do insumo na BOM."),
                "sugerido_mrp": st.column_config.NumberColumn("Sugerido pelo MRP", disabled=True,
                                                              help="Só para comparar; não entra no cálculo."),
                "data_entrega": st.column_config.DateColumn("Data de entrega", format="DD/MM/YYYY"),
                "entrega_em": st.column_config.SelectboxColumn("Entrega em", options=locais_po,
                                                               help="Terceiro que vai receber. Vazio = vale para qualquer."),
                "fornecedor": st.column_config.TextColumn("Fornecedor", width="medium"),
            })
        st.session_state["po_last"] = po_ed.copy()
        po_df = _po_para_calculo(po_ed)
        if len(po_df):
            st.caption(f"{len(po_df)} pedidos válidos · {int(po_df['data'].isna().sum())} sem data de entrega "
                       f"(sem data, o material não entra na programação).")
            buf = io.BytesIO()
            _po_para_excel(po_ed, buf)
            st.download_button("Baixar pedidos criados (Excel)", buf.getvalue(),
                               f"pedidos_compra_criados_{hoje:%Y%m%d}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # MRP final: ERP + pedidos criados aqui (viram "a caminho" com a data de entrega informada)
    kw["pedidos"] = _somar_pedidos(dados.pedidos, po_df)
    mrp = calcular_mrp(sim_compra, **kw) if tem_bom else pd.DataFrame()
    cen = comparar_cenarios(skus_edit.drop(columns=["ajuste_cx"]), dias_cen, kw, mult, desc_res) if tem_bom \
        else pd.DataFrame()

    if len(mrp):
        def _obs(r):
            o = []
            if r.sem_estoque_cadastrado:
                o.append("sem estoque cadastrado (tratado como 0)")
            if r.lead_time_padrao:
                o.append("lead time padrão")
            if r.sem_preco:
                o.append("sem preço")
            if r.sem_fornecedor:
                o.append("sem fornecedor")
            if r.transito_fora_prazo > 0:
                o.append(f"{theme.fmt_int(r.transito_fora_prazo)} em trânsito chega após a data")
            return "; ".join(o)
        mrp["obs"] = mrp.apply(_obs, axis=1)

    a_produzir = sim[sim["un_a_produzir"].fillna(0) > 0]
    estouros = ocup[ocup["status"] == "ESTOURA CAPACIDADE"]
    a_comprar = mrp[mrp["qtd_comprar"] > 0] if len(mrp) else mrp

    # ======================= Visão geral =======================
    with T["Visão geral"]:
        for _, r in estouros.iterrows():
            st.markdown(f'<div class="bz-alerta"><b>Capacidade estourada em {r["fabrica"]}:</b> '
                        f'{theme.fmt_int(r["ocupacao"])} posições de pallet para uma capacidade de '
                        f'{theme.fmt_int(r["capacidade_posicoes"])} (excede em {theme.fmt_int(-r["saldo"])}).</div>',
                        unsafe_allow_html=True)
        if len(mrp) and mrp["lead_time_padrao"].all():
            st.markdown(f'<div class="bz-alerta"><b>Nenhum fornecedor tem lead time cadastrado.</b> Todos usam '
                        f'{lt_padrao} dias — datas limite e o status ATRASADO são provisórios. '
                        f'Preencha na aba <i>Parâmetros</i>.</div>', unsafe_allow_html=True)
        if len(a_comprar):
            n_sem = int(a_comprar["sem_estoque_cadastrado"].sum())
            if n_sem:
                st.markdown(f'<div class="bz-alerta"><b>{n_sem} insumos a comprar não têm linha na planilha de '
                            f'estoque</b> (tratados como estoque zero). A compra pode estar superestimada — '
                            f'veja <i>Qualidade dos dados</i>.</div>', unsafe_allow_html=True)
        if mes_sel and sem_meta_mes:
            st.markdown(f'<div class="bz-alerta">Sem meta de {mes_sel.title()} para: {", ".join(sem_meta_mes)} '
                        f'(mantida a meta da planilha de reposição).</div>', unsafe_allow_html=True)

        k = st.columns(5)
        k[0].metric("SKUs a produzir", f"{len(a_produzir)} de {len(sim)}")
        k[1].metric("Unidades a produzir", theme.fmt_int(sim["un_a_produzir"].sum()))
        k[2].metric("Caixas", theme.fmt_int(sim["cx_mult_pallet"].sum()))
        k[3].metric("Pallets", theme.fmt_num(sim["pallets"].sum(), 1))
        k[4].metric("Compra de insumos (est.)", theme.fmt_brl(mrp["valor_estimado"].sum()) if len(mrp) else "—")
        k = st.columns(5)
        k[0].metric("Dias de cobertura alvo", f"{dias} dias")
        k[1].metric("SKUs com cobertura < 10 dias", int((sim["cob_atual_dias"] < 10).sum()))
        k[2].metric("Insumos a comprar", len(a_comprar) if len(mrp) else "—")
        k[3].metric("Insumos atrasados", int((mrp["status"] == "ATRASADO").sum()) if len(mrp) else "—")
        k[4].metric("Fábricas acima da capacidade", len(estouros))

        g1, g2 = st.columns(2)
        with g1:
            grafico_faixas(sim["faixa"], "Cobertura hoje")
        with g2:
            fin = pd.Series([faixa_cobertura(c, not pd.isna(c)) for c in sim["cob_final_dias"]])
            grafico_faixas(fin, f"Cobertura após produzir ({dias} dias alvo)")

        st.subheader("Cenários lado a lado")
        if len(cen):
            c1, c2 = st.columns(2)
            with c1:
                grafico_barra(cen, "dias_alvo", "pallets", "Pallets a produzir", theme.VERDE_ESCURO, ",.0f")
            with c2:
                grafico_barra(cen, "dias_alvo", "valor_compras", "Compra estimada de insumos (R$)", theme.VERDE_CLARO,
                              ",.0f")
            st.dataframe(cen, hide_index=True, width="stretch", column_config={
                "dias_alvo": "Dias alvo", "skus_a_produzir": "SKUs a produzir",
                "unidades": st.column_config.NumberColumn("Unidades", format="%.0f"),
                "caixas": st.column_config.NumberColumn("Caixas", format="%.0f"),
                "pallets": st.column_config.NumberColumn("Pallets", format="%.1f"),
                "posicoes_pallet": st.column_config.NumberColumn("Posições de pallet", format="%.0f"),
                "insumos_a_comprar": "Insumos a comprar", "insumos_atrasados": "Insumos atrasados",
                "valor_compras": st.column_config.NumberColumn("Compras (R$)", format="R$ %.2f"),
            })
            st.caption("Os cenários usam as mesmas premissas da barra lateral, sem os 'Forçar caixas' manuais.")
        else:
            st.info("Cenários dependem da estrutura de produtos (BOM) para calcular o MRP.")

        resumo = {
            "Gerado em": datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M"),
            "Dias de cobertura alvo": str(dias), "Meta de vendas": sel_meta,
            "Ajuste na demanda": f"{ajuste:+d}%", "Início planejado da produção": data_nec.strftime("%d/%m/%Y"),
            "Fábricas no MRP": ", ".join(fab_mrp), "Lead time padrão (dias)": str(lt_padrao),
            "Arquivos": ", ".join(arquivos),
        }
        xls = exportar_excel(sim, mrp, cen, ocup, dados.avisos, lead_df, cap_df, resumo)
        st.download_button("Baixar resultado completo (Excel)", xls, f"planejamento_mrp_pcp_{dias}d_{hoje:%Y%m%d}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")

    # ======================= Simulação PCP =======================
    with T["Simulação PCP"]:
        f1, f2, f3, f4 = st.columns([2, 3, 2, 3])
        fsel = f1.multiselect("Fábrica", fabs, default=fabs, key="pcp_fab")
        fxsel = f2.multiselect("Faixa de cobertura hoje", FAIXAS_ORDEM, default=FAIXAS_ORDEM, key="pcp_faixa")
        so_prod = f3.checkbox("Só quem produz", value=False)
        busca = f4.text_input("Buscar SKU / descrição", key="pcp_busca")
        v = sim[sim["fabrica"].isin(fsel) & sim["faixa"].isin(fxsel)]
        if so_prod:
            v = v[v["un_a_produzir"].fillna(0) > 0]
        if busca:
            b = busca.lower()
            v = v[v["sku"].str.lower().str.contains(b) | v["descricao"].str.lower().str.contains(b)]
        v = v.sort_values("cob_atual_dias", na_position="last")
        cols = ["sku", "descricao", "fabrica", "faixa", "estoque_un", "em_producao", "meta_mensal", "cob_atual_dias",
                "cob_pos_prod_dias", "nec_un", "cx_mult_pallet", "un_a_produzir", "pallets", "cob_final_dias",
                "pos_producao", "alerta_dados"]

        def _cor_faixa(x):
            return f"background-color:{theme.COR_FAIXA.get(x, '#FFFFFF')}33"

        st.dataframe(
            v[cols].style.map(_cor_faixa, subset=["faixa"]), hide_index=True, width="stretch", height=560,
            column_config={
                "sku": "SKU", "descricao": st.column_config.TextColumn("Descrição", width="large"), "fabrica": "Fábrica",
                "faixa": "Faixa hoje",
                "estoque_un": st.column_config.NumberColumn("Estoque (un.)", format="%.0f"),
                "em_producao": st.column_config.NumberColumn("Em produção", format="%.0f"),
                "meta_mensal": st.column_config.NumberColumn("Meta mês", format="%.0f"),
                "cob_atual_dias": st.column_config.NumberColumn("Cobertura hoje (d)", format="%.1f"),
                "cob_pos_prod_dias": st.column_config.NumberColumn("Cobertura pós OP (d)", format="%.1f"),
                "nec_un": st.column_config.NumberColumn("Necessário (un.)", format="%.0f"),
                "cx_mult_pallet": st.column_config.NumberColumn("Caixas (múlt. pallet)", format="%.0f"),
                "un_a_produzir": st.column_config.NumberColumn("A produzir (un.)", format="%.0f"),
                "pallets": st.column_config.NumberColumn("Pallets", format="%.1f"),
                "cob_final_dias": st.column_config.NumberColumn("Cobertura final (d)", format="%.0f"),
                "pos_producao": st.column_config.NumberColumn("Posições", format="%.0f"),
                "alerta_dados": "Alertas de dados",
            })
        st.caption(f"{len(v)} SKUs · {theme.fmt_int(v['un_a_produzir'].sum())} unidades · "
                   f"{theme.fmt_int(v['cx_mult_pallet'].sum())} caixas · {theme.fmt_num(v['pallets'].sum(), 1)} pallets")

    # ======================= Pallets por fábrica =======================
    with T["Pallets por fábrica"]:
        st.caption(f"Base da ocupação: **{BASES_OCUPACAO[base_oc]}**. Cada SKU fecha a sua própria posição "
                   f"(ROUNDUP de unidades ÷ unidades por pallet).")
        if len(ocup):
            for _, r in estouros.iterrows():
                st.error(f"{r['fabrica']}: {theme.fmt_int(r['ocupacao'])} posições > capacidade "
                         f"{theme.fmt_int(r['capacidade_posicoes'])} — excede em {theme.fmt_int(-r['saldo'])}.")
            for _, r in ocup[ocup["status"].str.startswith("SEM CAPACIDADE")].iterrows():
                st.warning(f"{r['fabrica']}: capacidade não informada (aba Parâmetros).")
            grafico_ocupacao(ocup.dropna(subset=["capacidade_posicoes"]))
            st.dataframe(ocup, hide_index=True, width="stretch", column_config={
                "fabrica": "Fábrica",
                "pos_estoque": st.column_config.NumberColumn("Posições estoque", format="%.0f"),
                "pos_em_producao": st.column_config.NumberColumn("Posições em produção", format="%.0f"),
                "pos_producao": st.column_config.NumberColumn("Posições simuladas", format="%.0f"),
                "skus_a_produzir": "SKUs a produzir",
                "ocupacao": st.column_config.NumberColumn("Ocupação", format="%.0f"),
                "capacidade_posicoes": st.column_config.NumberColumn("Capacidade", format="%.0f"),
                "saldo": st.column_config.NumberColumn("Saldo", format="%.0f"),
                "status": "Status", "uso_pct": st.column_config.ProgressColumn("Uso", format="percent", min_value=0,
                                                                              max_value=1.2),
            })
            fb = st.selectbox("Detalhar SKUs da fábrica", fabs)
            dd = sim[(sim["fabrica"] == fb) & (sim["pos_producao"].fillna(0) > 0)].sort_values("pos_producao",
                                                                                              ascending=False)
            st.dataframe(dd[["sku", "descricao", "un_a_produzir", "cx_mult_pallet", "pallets", "pos_producao"]],
                         hide_index=True, width="stretch")

    # ======================= Programação do terceiro =======================
    with T["Programação do terceiro"]:
        terceiros = sorted(dados.skus["fabrica"].unique())
        locais_com_saldo = set(dados.estoque["local"]) if dados.estoque is not None else set()
        pad = next((f for f in terceiros if f in locais_com_saldo), terceiros[0])
        c1, c2, c3 = st.columns([3, 2, 3])
        fab_t = c1.radio("Terceiro / fábrica", terceiros, index=terceiros.index(pad), horizontal=True, key="prog_fab")
        dias_prog = int(c2.number_input("Horizonte da programação (dias)", 7, 365, int(dias), step=5, key="prog_dias",
                                        help="Quantos dias de cobertura esta programação deve cobrir (ex.: 30)."))
        dias_apos = int(c3.number_input("Dias entre a chegada do material e o início da produção", 0, 30, 0, key="prog_apos"))
        pallet_f = st.checkbox("Produzir só pallets fechados (arredonda para baixo ao pallet quando o material não cobre tudo)",
                               value=True, key="prog_pallet",
                               help="Sempre em caixas fechadas. Com esta opção, também em múltiplos de pallet.")
        est_loc = dados.estoque[dados.estoque["local"] == fab_t] if dados.estoque is not None else pd.DataFrame()
        if est_loc.empty:
            st.warning(f"**Sem saldo de material de {fab_t}.** Envie o arquivo de estoque desse terceiro (aba com colunas "
                       f"'SKU Componente' e 'Estoque', com {fab_t} no nome) — sem ele não dá para dizer o que ele já "
                       f"consegue produzir.")
        elif not tem_bom:
            st.error("Estrutura de produtos (BOM) não encontrada.")
        else:
            sim_prog = sim if dias_prog == dias else simular_reposicao(skus_edit.drop(columns=["ajuste_cx"]), dias_prog,
                                                                         mult, desc_res)
            lista = sim_prog[(sim_prog["fabrica"] == fab_t) & (sim_prog["un_a_produzir"].fillna(0) > 0)].sort_values(
                ["cob_atual_dias", "sku"], na_position="last")
            if lista.empty:
                st.success(f"Nenhum SKU de {fab_t} precisa produzir para {dias_prog} dias.")
            else:
                dt_col = est_loc["data_saldo"].dropna() if "data_saldo" in est_loc.columns else []
                saldo_padrao = pd.Timestamp(max(dt_col)).date() if len(dt_col) else _data_do_saldo(
                    dados.abas_usadas.get("estoque_insumos", ""), hoje)
                data_saldo = st.date_input("Data do saldo de material informado", value=saldo_padrao,
                                           format="DD/MM/YYYY", key=f"prog_saldo_{fab_t}")
                idade = (hoje - data_saldo).days
                if idade > 7:
                    st.markdown(f'<div class="bz-alerta">O saldo de material tem <b>{idade} dias</b>. Peça o estoque '
                                f'atualizado ao terceiro antes de enviar a programação.</div>', unsafe_allow_html=True)

                with st.expander("Prioridade dos itens (edite o número para mudar a ordem de produção)", expanded=False):
                    pri_usr = lista["sku"].map(dados.skus.set_index("sku")["prioridade"])      # coluna prioridade da planilha
                    if pri_usr.notna().any():
                        lista = lista.assign(_p=pri_usr.fillna(1e9).values).sort_values(["_p", "cob_atual_dias", "sku"])
                    pri_base = pd.DataFrame({"prioridade": range(1, len(lista) + 1), "sku": lista["sku"].values,
                                             "descricao": lista["descricao"].values,
                                             "cobertura_hoje": lista["cob_atual_dias"].round(1).values})
                    pri_ed = st.data_editor(
                        pri_base, key=f"pri_{sig}_{fab_t}_{dias_prog}_{mes_sel}_{rst}", hide_index=True, width="stretch",
                        disabled=["sku", "descricao", "cobertura_hoje"],
                        column_config={"prioridade": st.column_config.NumberColumn("Prioridade (1 = primeiro)", min_value=1, step=1),
                                       "sku": "SKU", "descricao": st.column_config.TextColumn("Descrição", width="large"),
                                       "cobertura_hoje": st.column_config.NumberColumn("Cobertura hoje (dias)", format="%.1f")})
                    st.caption("Padrão: menor cobertura primeiro. O material em estoque é compartilhado, então quem vem "
                               "primeiro consome primeiro.")
                prio = pri_ed.set_index("sku")["prioridade"]
                sim_t = sim_prog.copy()
                sim_t["prioridade"] = sim_t["sku"].map(prio)

                entradas = po_df[po_df["entrega_em"].isin(["", fab_t])][["comp_sku", "qtd", "data"]] if len(po_df) else None
                por_sku, ondas, faltas = cronograma_terceiro(sim_t, dados.bom, est_loc, entradas, fab_t,
                                                             pd.Timestamp(hoje), dias_apos, pallet_f)

                n = len(por_sku)
                cont = por_sku["status"].value_counts()
                k = st.columns(5)
                k[0].metric("SKUs no plano", n)
                k[1].metric("Produz tudo agora", int(cont.get("PRODUZ TUDO AGORA", 0)))
                k[2].metric("Produz parte agora", int(cont.get("PARTE AGORA, RESTO COM PEDIDO", 0) + cont.get("PARTE SEM PEDIDO", 0)))
                k[3].metric("Aguarda material (com pedido)", int(cont.get("AGUARDA MATERIAL", 0)))
                k[4].metric("Sem pedido que cubra", int(cont.get("SEM PEDIDO QUE COBRE", 0) + cont.get("PARTE SEM PEDIDO", 0)))
                k = st.columns(4)
                k[0].metric("Caixas do plano", theme.fmt_int(por_sku["plano_cx"].sum()))
                k[1].metric("Pode produzir agora (cx)", theme.fmt_int(por_sku["agora_cx"].sum()))
                k[2].metric("Libera com os pedidos (cx)", theme.fmt_int(por_sku["depois_cx"].sum()))
                k[3].metric("Ainda sem pedido (cx)", theme.fmt_int(por_sku["sem_pedido_cx"].sum()))

                COR_ST = {"PRODUZ TUDO AGORA": "background-color:#DDEBE2;color:#0B592F;font-weight:600",
                          "PARTE AGORA, RESTO COM PEDIDO": "background-color:#E2EBB8;color:#3D4A00;font-weight:600",
                          "AGUARDA MATERIAL": "background-color:#FBE9B5;color:#6B4B00;font-weight:600",
                          "PARTE SEM PEDIDO": "background-color:#FBE9B5;color:#6B4B00;font-weight:600",
                          "SEM PEDIDO QUE COBRE": "background-color:#F6D5C6;color:#7A2E0E;font-weight:600",
                          "SEM ESTRUTURA": "background-color:#D9D9D9;color:#3D3D3D"}
                ordem_pri = {s_: i + 1 for i, s_ in enumerate(por_sku["sku"])}
                tab = por_sku.copy()
                tab.insert(0, "prioridade", tab["sku"].map(ordem_pri))
                tab["agora_pallets"] = tab["agora_cx"] / tab["cx_por_pallet"]
                st.dataframe(
                    tab[["prioridade", "sku", "descricao", "cob_atual_dias", "plano_cx", "agora_cx", "agora_pallets",
                         "depois_cx", "sem_pedido_cx", "primeira_data", "status"]].style.map(lambda x: COR_ST.get(x, ""), subset=["status"]),
                    hide_index=True, width="stretch", height=min(40 + 35 * n, 560),
                    column_config={
                        "prioridade": st.column_config.NumberColumn("Prior.", format="%d", width="small"),
                        "sku": "SKU", "descricao": st.column_config.TextColumn("Descrição", width="large"),
                        "cob_atual_dias": st.column_config.NumberColumn("Cobertura hoje (d)", format="%.1f"),
                        "plano_cx": st.column_config.NumberColumn("Plano (cx)", format="%.0f"),
                        "agora_cx": st.column_config.NumberColumn("Produz agora (cx)", format="%.0f"),
                        "agora_pallets": st.column_config.NumberColumn("Produz agora (pallets)", format="%.0f"),
                        "depois_cx": st.column_config.NumberColumn("Libera c/ pedidos (cx)", format="%.0f"),
                        "sem_pedido_cx": st.column_config.NumberColumn("Sem pedido (cx)", format="%.0f"),
                        "primeira_data": st.column_config.DateColumn("Próxima liberação", format="DD/MM/YYYY"),
                        "status": "Situação"})

                g1, g2 = st.columns(2)
                with g1:
                    st.markdown("**Cronograma: o que é produzido em cada data**")
                    if len(ondas):
                        cr = ondas.groupby(["data", "sku", "descricao"], as_index=False).agg(
                            cx=("cx", "sum"), pallets=("pallets", "sum"), un=("un", "sum"))
                        cr = cr.sort_values(["data", "sku"])
                        st.dataframe(cr, hide_index=True, width="stretch", column_config={
                            "data": st.column_config.DateColumn("A partir de", format="DD/MM/YYYY"), "sku": "SKU",
                            "descricao": st.column_config.TextColumn("Descrição", width="large"),
                            "cx": st.column_config.NumberColumn("Caixas", format="%.0f"),
                            "pallets": st.column_config.NumberColumn("Pallets", format="%.0f"),
                            "un": st.column_config.NumberColumn("Unidades", format="%.0f")})
                    else:
                        st.caption("Nada pode ser produzido com o saldo e os pedidos atuais.")
                with g2:
                    st.markdown("**Material que ainda falta (sem pedido)**")
                    if len(faltas):
                        st.dataframe(faltas, hide_index=True, width="stretch", column_config={
                            "comp_sku": "Insumo", "descricao": st.column_config.TextColumn("Descrição", width="large"),
                            "unidade": "Un.", "falta": st.column_config.NumberColumn("Falta", format="%.1f")})
                        st.caption("Cadastre os pedidos na aba *Pedidos de compra criados* e a simulação se atualiza.")
                    else:
                        st.success("Todo o plano tem material em estoque ou com pedido.")

                st.markdown("---")
                obs_txt = st.text_area("Observações para o terceiro (vão no topo do PDF)", key=f"prog_obs_{fab_t}",
                                       placeholder="Ex.: Priorizar os itens na ordem indicada. Favor confirmar o saldo atual.")
                meta_pdf = {"Horizonte do plano": f"{dias_prog} dias", "Saldo de material": f"{data_saldo:%d/%m/%Y}"}
                pdf_bytes = gerar_pdf_programacao(fab_t, por_sku, ondas, faltas, pd.Timestamp(hoje), meta_pdf, obs_txt)
                d1, d2 = st.columns(2)
                d1.download_button(f"Baixar PDF para {fab_t}", pdf_bytes, f"programacao_{fab_t}_{hoje:%Y%m%d}.pdf",
                                   mime="application/pdf", type="primary", width="stretch")
                xb = io.BytesIO()
                with pd.ExcelWriter(xb, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
                    tab.rename(columns={"prioridade": "Prioridade", "sku": "SKU", "descricao": "Descrição",
                                        "cob_atual_dias": "Cobertura hoje (d)", "plano_cx": "Plano (cx)",
                                        "agora_cx": "Produz agora (cx)", "depois_cx": "Libera c/ pedidos (cx)",
                                        "sem_pedido_cx": "Sem pedido (cx)", "ultima_data": "Conclusão",
                                        "primeira_data": "Próxima liberação", "status": "Situação"}).to_excel(
                        w, index=False, sheet_name="Programação")
                    ondas.drop(columns=["onda"]).to_excel(w, index=False, sheet_name="Cronograma")
                    faltas.to_excel(w, index=False, sheet_name="Material em falta")
                d2.download_button(f"Baixar Excel da programação ({fab_t})", xb.getvalue(),
                                   f"programacao_{fab_t}_{hoje:%Y%m%d}.xlsx",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")

    # ======================= MRP =======================
    with T["MRP Compras"]:
        if not tem_bom:
            st.error("Estrutura de produtos (BOM) não encontrada: não é possível calcular o MRP.")
        elif mrp.empty:
            st.info("Nenhum SKU com produção simulada nas fábricas selecionadas.")
        else:
            st.caption(f"Insumos necessários para o simulado de **{dias_compra} dias** (produção: {dias}), com início em "
                       f"**{data_nec:%d/%m/%Y}**. Necessidade líquida = consumo bruto − (estoque + pedidos em trânsito).")
            m1, m2, m3, m4 = st.columns([3, 3, 3, 3])
            st_opts = ["ATRASADO", "PEDIR ESTA SEMANA", "PROGRAMAR", "COBERTO"]
            st_sel = m1.multiselect("Status", st_opts, default=st_opts[:3], key="mrp_status")
            forn_opts = sorted(mrp["fornecedor"].unique())
            forn_sel = m2.multiselect("Fornecedor", forn_opts, default=[], key="mrp_forn",
                                      placeholder="Todos")
            ocultar = m3.checkbox("Ocultar sem estoque cadastrado", value=False)
            busca_m = m4.text_input("Buscar insumo", key="mrp_busca")
            v = mrp[mrp["status"].isin(st_sel)]
            if forn_sel:
                v = v[v["fornecedor"].isin(forn_sel)]
            if ocultar:
                v = v[~v["sem_estoque_cadastrado"]]
            if busca_m:
                b = busca_m.lower()
                v = v[v["comp_sku"].str.lower().str.contains(b) | v["descricao"].str.lower().str.contains(b)]

            def _cor_status(x):
                return theme.COR_STATUS_MRP.get(x, "")

            cols = ["status", "comp_sku", "descricao", "fornecedor", "unidade", "consumo_bruto", "estoque",
                    "transito_considerado", "necessidade_liquida", "qtd_sugerida", "preco", "valor_estimado",
                    "lead_time_dias", "data_limite_pedido", "dias_ate_limite", "obs"]
            st.dataframe(
                v[cols].style.map(_cor_status, subset=["status"]), hide_index=True, width="stretch", height=520,
                column_config={
                    "status": "Status", "comp_sku": "Insumo", "descricao": st.column_config.TextColumn("Descrição", width="large"),
                    "fornecedor": st.column_config.TextColumn("Fornecedor", width="medium"), "unidade": "Un.",
                    "consumo_bruto": st.column_config.NumberColumn("Consumo bruto", format="%.2f"),
                    "estoque": st.column_config.NumberColumn("Estoque", format="%.2f"),
                    "transito_considerado": st.column_config.NumberColumn("Em trânsito", format="%.2f"),
                    "necessidade_liquida": st.column_config.NumberColumn("Necessidade líquida", format="%.2f"),
                    "qtd_sugerida": st.column_config.NumberColumn("Qtd a comprar", format="%.2f"),
                    "preco": st.column_config.NumberColumn("Preço unit.", format="R$ %.4f"),
                    "valor_estimado": st.column_config.NumberColumn("Valor estimado", format="R$ %.2f"),
                    "lead_time_dias": st.column_config.NumberColumn("Lead time (d)", format="%d"),
                    "data_limite_pedido": st.column_config.DateColumn("Pedir até", format="DD/MM/YYYY"),
                    "dias_ate_limite": st.column_config.NumberColumn("Dias p/ limite", format="%d"),
                    "obs": st.column_config.TextColumn("Observações", width="large"),
                })
            st.caption(f"{len(v)} insumos · valor estimado {theme.fmt_brl(v['valor_estimado'].sum())}")

            with st.expander("Resumo por fornecedor (para emitir os pedidos)"):
                rf = (mrp[mrp["qtd_comprar"] > 0].groupby("fornecedor")
                      .agg(itens=("comp_sku", "count"), valor=("valor_estimado", "sum"),
                           pedir_ate=("data_limite_pedido", "min"),
                           atrasados=("status", lambda s: int((s == "ATRASADO").sum())))
                      .reset_index().sort_values("pedir_ate"))
                st.dataframe(rf, hide_index=True, width="stretch", column_config={
                    "fornecedor": "Fornecedor", "itens": "Itens",
                    "valor": st.column_config.NumberColumn("Valor estimado", format="R$ %.2f"),
                    "pedir_ate": st.column_config.DateColumn("Pedir até (mais urgente)", format="DD/MM/YYYY"),
                    "atrasados": "Itens atrasados"})

    # ======================= Consumo por insumo =======================
    with T["Consumo por insumo"]:
        if not tem_bom or mrp.empty:
            st.info("Sem MRP calculado.")
        else:
            det = consumo_detalhado(sim[sim["fabrica"].isin(fab_mrp)], dados.bom)
            desc_map = mrp.set_index("comp_sku")["descricao"].to_dict()
            ins = st.selectbox("Insumo", sorted(mrp["comp_sku"]), format_func=lambda s: f"{s} — {desc_map.get(s, '')}")
            d1 = det[det["comp_sku"] == ins].rename(columns={"kit_sku": "sku"}).sort_values("consumo", ascending=False)
            r = mrp[mrp["comp_sku"] == ins].iloc[0]
            c = st.columns(4)
            c[0].metric("Consumo bruto", theme.fmt_num(r["consumo_bruto"], 2))
            c[1].metric("Estoque", theme.fmt_num(r["estoque"], 2))
            c[2].metric("Em trânsito (considerado)", theme.fmt_num(r["transito_considerado"], 2))
            c[3].metric("Necessidade líquida", theme.fmt_num(r["necessidade_liquida"], 2))
            st.markdown("**SKUs que consomem este insumo**")
            st.dataframe(d1[["sku", "un_a_produzir", "qtd", "un", "consumo"]], hide_index=True, width="stretch",
                         column_config={"sku": "SKU", "un_a_produzir": st.column_config.NumberColumn("A produzir", format="%.0f"),
                                        "qtd": st.column_config.NumberColumn("Qtd por unidade", format="%.4f"),
                                        "un": "Un.", "consumo": st.column_config.NumberColumn("Consumo", format="%.2f")})
            if dados.pedidos is not None and "comp_sku" in dados.pedidos.columns:
                pp = dados.pedidos[(dados.pedidos["comp_sku"] == ins)]
                st.markdown("**Pedidos em trânsito deste insumo**")
                if len(pp):
                    st.dataframe(pp[["oc", "data", "prevista", "fornecedor", "qtd"]], hide_index=True, width="stretch",
                                 column_config={"oc": "OC", "data": st.column_config.DateColumn("Emissão", format="DD/MM/YYYY"),
                                                "prevista": st.column_config.DateColumn("Previsão", format="DD/MM/YYYY"),
                                                "fornecedor": "Fornecedor", "qtd": "Quantidade"})
                else:
                    st.caption("Nenhum pedido em trânsito.")

    # ======================= Qualidade dos dados =======================
    with T["Qualidade dos dados"]:
        st.subheader("Abas utilizadas")
        st.dataframe(pd.DataFrame([(ROTULO_PAPEL.get(p, "Pedidos de compra criados"), a) for p, a in dados.abas_usadas.items()],
                                  columns=["Dado", "Origem (arquivo › aba)"]), hide_index=True, width="stretch")
        st.subheader("Pontos de atenção")
        ordem = {"erro": 0, "aviso": 1, "info": 2}
        for nivel, msg in sorted(dados.avisos, key=lambda x: ordem.get(x[0], 9)):
            css = "bz-alerta" if nivel in ("erro", "aviso") else "bz-ok"
            rot = {"erro": "Erro", "aviso": "Atenção", "info": "Info"}[nivel]
            st.markdown(f'<div class="{css}"><b>{rot}:</b> {msg}</div>', unsafe_allow_html=True)
        if not dados.avisos:
            st.success("Nenhum ponto de atenção.")
        if dados.pedidos is not None and "tipo" in dados.pedidos.columns:
            nm = dados.pedidos[dados.pedidos["tipo"] == "nao_mapeado"]
            if len(nm):
                st.subheader("Pedidos que não casam com a BOM")
                st.dataframe(nm[["oc", "fornecedor", "item", "qtd"]], hide_index=True, width="stretch")

    st.markdown('<div class="bz-rodape">Planejador MRP/PCP · Bioz Green — Pelo Planeta · '
                'compras@biozgreen.eco.br</div>', unsafe_allow_html=True)


main()
